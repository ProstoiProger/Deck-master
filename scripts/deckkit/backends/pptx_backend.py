"""Draw scenes into an editable .pptx (real text boxes, shapes, speaker notes). Needs python-pptx."""
import re
from pathlib import Path


def _rgb(h):
    from pptx.dml.color import RGBColor
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_alpha(fill_fmt, alpha):
    """python-pptx has no transparency API; add <a:alpha> to the solid fill color."""
    from lxml import etree
    srgb = fill_fmt._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    if srgb is not None:
        a = etree.SubElement(srgb, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
        a.set("val", str(int(alpha * 100000)))


def _runs(text):
    """Split `**bold**` markup into (text, bold) runs."""
    out, state = [], False
    for i, part in enumerate(re.split(r"\*\*", text)):
        if i:
            state = not state
        if part:
            out.append((part, state))
    return out


def build_pptx(scenes, theme, deck, path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    from ..core.scene import H, W, Image, Line, Shape, Text

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

    for sc in scenes:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(sc.bg)
        for it in sc.items:
            if isinstance(it, Shape):
                kind = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL}[it.kind]
                sh = slide.shapes.add_shape(kind, Inches(it.x), Inches(it.y), Inches(it.w), Inches(it.h))
                if it.kind == "round":
                    sh.adjustments[0] = min(0.5, it.radius / max(min(it.w, it.h), 0.01))
                if it.fill:
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = _rgb(it.fill)
                    if it.alpha < 1:
                        _set_alpha(sh.fill, it.alpha)
                else:
                    sh.fill.background()
                if it.line:
                    sh.line.color.rgb = _rgb(it.line)
                    sh.line.width = Pt(it.line_w)
                else:
                    sh.line.fill.background()
                sh.shadow.inherit = False
                sh.text_frame.text = ""
            elif isinstance(it, Text):
                tb = slide.shapes.add_textbox(Inches(it.x), Inches(it.y), Inches(it.w), Inches(it.h))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = anchor[it.valign]
                for pi, line in enumerate(it.text.split("\n")):
                    p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
                    p.alignment = align[it.align]
                    p.line_spacing = it.line_h
                    for txt, b in _runs(line) or [("", False)]:
                        r = p.add_run()
                        r.text = txt
                        f = r.font
                        f.name = theme.pptx_font(it.role)
                        f.size = Pt(it.size)
                        f.bold = bool(it.bold or b)
                        f.color.rgb = _rgb(it.color)
            elif isinstance(it, Image):
                slide.shapes.add_picture(it.path, Inches(it.x), Inches(it.y), Inches(it.w), Inches(it.h))
            elif isinstance(it, Line):
                ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(it.x1), Inches(it.y1), Inches(it.x2), Inches(it.y2))
                ln.line.color.rgb = _rgb(it.color)
                ln.line.width = Pt(it.width)
        if sc.notes:
            slide.notes_slide.notes_text_frame.text = sc.notes

    cp = prs.core_properties
    cp.title = str(deck.meta.get("title", "")) or scenes[0].title
    cp.author = str(deck.meta.get("team", "")) or ""
    cp.subject = str(deck.meta.get("event", ""))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
